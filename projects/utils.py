import xml.etree.ElementTree as ET
from .models import Vulnerability, Project
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
import io

def parse_nmap_xml(xml_file, project_id):
    """
    Parses Nmap XML output and maps open ports/services to the database.
    """
    try:
        # Load the project instance
        project = Project.objects.get(id=project_id)
        
        # Parse the XML file
        tree = ET.parse(xml_file)
        root = tree.getroot()

        # Iterate through each 'host' in the Nmap scan
        for host in root.findall('host'):
            # Get the IP Address
            addr_node = host.find('address')
            ip_address = addr_node.get('addr') if addr_node is not None else "Unknown"

            # Iterate through each 'port' for the host
            for port in host.findall('.//port'):
                port_id = port.get('portid')
                state = port.find('state').get('state')

                # Only process open ports (Attack Surface)
                if state == 'open':
                    service_node = port.find('service')
                    service_name = service_node.get('name') if service_node is not None else "unknown"
                    product = service_node.get('product', '')
                    version = service_node.get('version', '')

                    # Create a technical finding in the database
                    Vulnerability.objects.create(
                        project=project,
                        title=f"Open Port Detected: {port_id}/{service_name}",
                        severity="Low", # Nmap findings are informational by default
                        ip_address=ip_address,
                        description=f"Port {port_id} is open running {product} {version}.",
                        remediation="Evaluate if this service is necessary. If not, close the port or restrict access via firewall/ACL."
                    )
        return True
    except Exception as e:
        print(f"Parsing Error: {e}")
        return False
    
def generate_assessment_report(project_id):
    """
    Generates a professional Word report for the given project.
    """
    project = Project.objects.get(id=project_id)
    vulnerabilities = project.vulnerabilities.all()
    
    doc = Document()
    
    # Report Title
    title = doc.add_heading(f'Security Assessment Report: {project.name}', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # Project Information Section
    doc.add_heading('1. Executive Summary', level=1)
    doc.add_paragraph(f"Client Name: {project.client_name}")
    doc.add_paragraph(f"Assessment Date: {project.created_at.strftime('%Y-%m-%d')}")
    doc.add_paragraph(f"This report details the findings identified during the security assessment of {project.name}.")
    
    # Findings Summary Table
    doc.add_heading('2. Findings Summary', level=1)
    table = doc.add_table(rows=1, cols=3)
    table.style = 'Table Grid'
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Vulnerability Title'
    hdr_cells[1].text = 'Severity'
    hdr_cells[2].text = 'IP Address'
    
    for vuln in vulnerabilities:
        row_cells = table.add_row().cells
        row_cells[0].text = vuln.title
        row_cells[1].text = vuln.severity
        row_cells[2].text = str(vuln.ip_address)

    # Detailed Findings Section
    doc.add_page_break()
    doc.add_heading('3. Technical Findings & Remediation', level=1)
    
    for vuln in vulnerabilities:
        doc.add_heading(vuln.title, level=2)
        doc.add_paragraph(f"Severity: {vuln.severity}")
        doc.add_paragraph(f"Affected Asset: {vuln.ip_address}")
        
        doc.add_heading('Description', level=3)
        doc.add_paragraph(vuln.description)
        
        doc.add_heading('Remediation Strategies', level=3)
        doc.add_paragraph(vuln.remediation)
        
        doc.add_paragraph("-" * 30)

    # Save the document to a byte stream
    file_stream = io.BytesIO()
    doc.save(file_stream)
    file_stream.seek(0)
    return file_stream